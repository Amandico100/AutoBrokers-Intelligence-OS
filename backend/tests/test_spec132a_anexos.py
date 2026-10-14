# -*- coding: utf-8 -*-
"""SPEC-132-A · A2 — OS ANEXOS DO CHAT, LIDOS POR BYTES (`chat_anexos.preparar`).

    BFF {attachments[]} → chat_anexos.preparar(company_id, anexos) → download do Storage (DUBLÊ) →
    imagem: JPEG ≤ 2048 em data URL + `describe_image` (o LLM do papel `visao` é DUBLÊ) · PDF: o leitor REAL da 130-B
    (o modelo do leitor é DUBLÊ) · xlsx/xls/csv: o leitor tabular REAL · docx/pptx: REAL · áudio: `AudioService` REAL
    com o cliente da API de transcrição DUBLÊ → tudo pela máscara REAL da 130-B.

Dublê SÓ na borda: o Storage, o LLM de visão, o modelo do leitor de PDF, a API de transcrição e o docling
(microserviço). A rede fica FECHADA (só o loopback do laço de eventos do Windows).

⛔ Dados fictícios: CPF sintético com DV válido, corretoras com UUID inventado, DDD 00.
Rodar: cd backend && python -m pytest tests/test_spec132a_anexos.py -q -p no:cacheprovider
"""
from __future__ import annotations

import asyncio
import io
import re
import socket
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

import app.services.audio_service as AU  # noqa: E402
import app.services.chat_anexos as CA  # noqa: E402
import app.services.policy_document_evidence_service as P  # noqa: E402
import app.services.vision_service as VS  # noqa: E402

EMPRESA_A = "aaaaaaaa-0000-4000-8000-000000132a01"
EMPRESA_B = "bbbbbbbb-0000-4000-8000-000000132a02"
DONO = "cccccccc-0000-4000-8000-000000132a03"


def _cpf_sintetico(base9: str) -> str:
    d = [int(c) for c in base9]
    for _ in range(2):
        s = sum(v * p for v, p in zip(d, range(len(d) + 1, 1, -1)))
        d.append(0 if s % 11 < 2 else 11 - s % 11)
    t = "".join(map(str, d))
    return f"{t[:3]}.{t[3:6]}.{t[6:9]}-{t[9:]}"


CPF = _cpf_sintetico("123456789")
CPF_DIGITOS = re.sub(r"\D", "", CPF)
#: um CPF que começa com ZERO: o Excel guarda como número e PERDE o zero (10 dígitos)
CPF_ZERO = _cpf_sintetico("012345678")
CPF_ZERO_NUMERO = int(re.sub(r"\D", "", CPF_ZERO))
TELEFONE = "(00) 99999-0001"
PII = (CPF, CPF_DIGITOS, CPF_ZERO, str(CPF_ZERO_NUMERO), re.sub(r"\D", "", CPF_ZERO), TELEFONE, "99999-0001",
       "15/03/1980")


def _sem_pii(texto: str) -> list:
    return [p for p in PII if p in texto]


def _path(empresa: str, bucket: str, ext: str, n: int = 1) -> str:
    return f"{empresa}/{DONO}/{bucket}/{n:08d}-2222-4333-8444-555555555555.{ext}"


# =====================================================================================================================
# os arquivos de teste, gerados aqui
# =====================================================================================================================
def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(["Nome", "CPF", "Nascimento", "Prêmio", "Observação", "Nº apólice"])
    ws.append(["Pessoa Ficticia Um", CPF, datetime(1980, 3, 15), 1234.56, f"titular CPF {CPF}", 900000123])
    ws.append(["Pessoa Ficticia Dois", CPF_ZERO_NUMERO, datetime(1975, 1, 2), 765.44, f"ligar {TELEFONE}", 900000124])
    ws2 = wb.create_sheet("Resumo")
    ws2.append(["Ramo", "Quantidade"])
    ws2.append(["Auto", 2])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_ponto_e_virgula() -> bytes:
    linhas = ["Segurado;Placa do carro;Valor;Obs", f"Pessoa Ficticia;ABC1D23;1.500,00;CPF {CPF}",
              "Outra Pessoa;XYZ9K87;2.500,50;renovar em março"]
    return ("\r\n".join(linhas) + "\r\n").encode("cp1252")


def _png(lado: int = 64) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGBA", (lado, lado // 2), (200, 30, 30, 128)).save(buf, format="PNG")
    return buf.getvalue()


def _heic() -> bytes:
    import pillow_heif
    from PIL import Image

    pillow_heif.register_heif_opener()
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (10, 120, 200)).save(buf, format="HEIF")
    return buf.getvalue()


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Proposta de renovação do seguro auto")
    d.add_paragraph(f"Cliente com CPF {CPF} pediu desconto")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    """Um .pptx mínimo (só o que o leitor lê: o XML dos slides dentro do pacote)."""
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    slide = (f'<p:sld {ns} xmlns:p="x"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>Carteira 2026</a:t></a:r></a:p>'
             f'<a:p><a:r><a:t>Contato CPF {CPF}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/slides/slide1.xml", slide)
    return buf.getvalue()


def _pdf_minimo(linhas) -> bytes:
    """Um PDF de UMA página escrito à mão (o mesmo molde do test_spec130b_anexos: o PyMuPDF é lento nesta máquina)."""
    def _esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    corpo = ["BT", "/F1 10 Tf", "14 TL", "60 780 Td"] + [f"({_esc(ln)}) Tj T*" for ln in linhas] + ["ET"]
    stream = "\n".join(corpo).encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 5 0 R >> >> "
            b"/Contents 4 0 R >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"]
    saida = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(saida))
        saida += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(saida)
    saida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        saida += b"%010d 00000 n \n" % off
    saida += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(saida)


def _pdf() -> bytes:
    return _pdf_minimo(["Ata da reuniao do condominio", f"Sindico: Pessoa Ficticia   CPF: {CPF}",
                        f"Telefone: {TELEFONE}", "Pauta: pintura da fachada e troca do portao"])


# =====================================================================================================================
# as bordas
# =====================================================================================================================
_CONNECT_REAL = socket.socket.connect
_CONNECT_EX_REAL = socket.socket.connect_ex


def _loopback(destino) -> bool:
    host = destino[0] if isinstance(destino, tuple) and destino else ""
    return str(host) in ("127.0.0.1", "::1", "localhost")


def _sem_rede(self, destino, *a, **k):
    if _loopback(destino):
        return _CONNECT_REAL(self, destino, *a, **k)
    raise AssertionError("a rede é fechada neste teste")


def _sem_rede_ex(self, destino, *a, **k):
    if _loopback(destino):
        return _CONNECT_EX_REAL(self, destino, *a, **k)
    raise AssertionError("a rede é fechada neste teste")


class _SanitizacaoForaDoAr:
    def _docling_parse(self, *_a, **_k):
        raise RuntimeError("docling fora do ar (teste)")


class _Storage:
    """O Supabase Storage: um dicionário (bucket, path) → bytes. Registra cada download."""

    def __init__(self):
        self.objetos = {}
        self.baixados = []

    def por(self, bucket, path, dados):
        self.objetos[(bucket, path)] = dados
        return {"bucket": bucket, "path": path}

    async def baixar(self, bucket, path):
        self.baixados.append((bucket, path))
        if (bucket, path) not in self.objetos:
            raise RuntimeError("Object not found")
        return self.objetos[(bucket, path)]


class _LlmDaVisao:
    """O modelo do papel `visao`: registra o que recebeu e devolve uma descrição COM um CPF (que a máscara tem de tirar)."""

    def __init__(self):
        self.recebido = []

    async def ainvoke(self, mensagens):
        self.recebido.append(mensagens)

        class _R:
            content = f"Foto de um para-choque amassado. Etiqueta com CPF {CPF} visível."

        return _R()


class _ApiDeTranscricao:
    def __init__(self, texto):
        self.texto = texto
        self.recebidos = []
        outer = self

        class _T:
            async def create(self, file=None, **k):
                outer.recebidos.append((getattr(file, "name", ""), k.get("model")))

                class _R:
                    text = outer.texto
                    usage = None

                return _R()

        class _A:
            transcriptions = _T()

        self.audio = _A()


@pytest.fixture(autouse=True)
def bordas(monkeypatch):
    monkeypatch.setattr(socket.socket, "connect", _sem_rede)
    monkeypatch.setattr(socket.socket, "connect_ex", _sem_rede_ex)
    from app.services import sanitization_service as S

    monkeypatch.setattr(S, "SanitizationService", _SanitizacaoForaDoAr)
    storage = _Storage()
    monkeypatch.setattr(CA, "_baixar_do_storage", storage.baixar)
    llm = _LlmDaVisao()
    monkeypatch.setattr(VS, "criar_llm_de_visao", lambda **_k: llm)
    api = _ApiDeTranscricao(f"Oi, é sobre o sinistro do carro, meu CPF é {CPF} e o telefone {TELEFONE}.")
    monkeypatch.setattr(CA, "_servico_de_audio", lambda: AU.AudioService("chave-de-teste", modelo="gpt-transcribe",
                                                                         cliente=api))
    monkeypatch.setattr(AU.AudioService, "_registrar_uso", lambda *a, **k: None)
    monkeypatch.setattr(AU, "prompt_de_dominio", lambda: "Atendimento de corretora de seguros no Brasil.")
    # o leitor REAL da 130-B, com o MODELO dele dublado (devolve vazio → o regex e o texto mascarado)
    original = P.ler_anexo_de_seguro

    async def _vazio(*_a, **_k):
        return None

    async def _leitor(file_bytes, file_name, *, company_id, orcamento_s=40.0, leitor_modelo=None):
        return await original(file_bytes, file_name, company_id=company_id, orcamento_s=orcamento_s,
                              leitor_modelo=_vazio)

    monkeypatch.setattr(P, "ler_anexo_de_seguro", _leitor)
    return {"storage": storage, "llm": llm, "api": api}


def _rodar(empresa, anexos, **k):
    return asyncio.run(CA.preparar(empresa, anexos, **k))


def _anexo(storage, empresa, bucket, ext, dados, nome, mime, n=1):
    ref = storage.por(bucket, _path(empresa, bucket, ext, n), dados)
    return {**ref, "nome": nome, "mime": mime, "bytes": len(dados)}


# =====================================================================================================================
# (a) o CPF da planilha NÃO chega cru ao texto
# =====================================================================================================================
def test_a_xlsx_le_abas_cabecalho_totais_e_nenhum_cpf_cru(bordas):
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "xlsx", _xlsx(), "carteira.xlsx",
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r = _rodar(EMPRESA_A, [a])
    assert r.avisos == [], r.avisos
    t = r.texto
    assert _sem_pii(t) == [], f"PII crua no texto da planilha: {_sem_pii(t)}"
    assert t.startswith("[DOCUMENTO ANEXADO: carteira.xlsx]") and t.rstrip().endswith("[FIM DO DOCUMENTO]"), t[:200]
    assert 'Aba "Clientes"' in t and 'Aba "Resumo"' in t and "2 linhas de dados × 6 colunas" in t, t[:600]
    assert "Colunas: Nome, CPF, Nascimento, Prêmio, Observação, Nº apólice" in t, t[:600]
    assert "Pessoa Ficticia Um" in t and "| 1.234,56 |" in t, t
    assert "- Prêmio: soma 2.000,00 · mín 765,44 · máx 1.234,56 (2 valores)" in t, t
    assert "Nº apólice: soma" not in t and "Quantidade: soma 2" in t, t
    assert "[CPF]" in t and "[DATA]" in t, t
    assert r.lidos == [{"nome": "carteira.xlsx", "mime": a["mime"], "bytes": a["bytes"], "path": a["path"],
                        "bucket": "chat-docs", "tipo": "xlsx", "lido": True}], r.lidos


def test_a_controle_sem_a_mascara_o_cpf_da_planilha_passaria(bordas, monkeypatch):
    """O guarda (a) CONSEGUE falhar: com a máscara desligada (identidade), o CPF da coluna "Observação" sai cru."""
    monkeypatch.setattr(CA, "_mascarar", lambda texto: texto)
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "xlsx", _xlsx(), "carteira.xlsx",
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert CPF in _rodar(EMPRESA_A, [a]).texto


def test_a_mascara_indisponivel_e_fail_closed(bordas, monkeypatch):
    monkeypatch.setattr(CA, "_mascarar", lambda texto: None)
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "csv", _csv_ponto_e_virgula(), "lista.csv", "text/csv")
    r = _rodar(EMPRESA_A, [a])
    assert r.texto == "" and r.avisos and "com segurança" in r.avisos[0], (r.texto, r.avisos)


def test_csv_com_ponto_e_virgula_e_cp1252(bordas):
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "csv", _csv_ponto_e_virgula(), "lista.csv", "application/vnd.ms-excel")
    r = _rodar(EMPRESA_A, [a])
    t = r.texto
    assert r.avisos == [] and _sem_pii(t) == [], (r.avisos, _sem_pii(t))
    assert "2 linhas de dados × 4 colunas" in t and "renovar em março" in t, t
    assert "- Valor: soma 4.000,50 · mín 1.500,00 · máx 2.500,50 (2 valores)" in t, t
    assert "ABC1D23" not in t and "[PLACA]" in t, t
    assert r.lidos[0]["tipo"] == "csv", "o .csv que o Windows manda como vnd.ms-excel é lido como CSV"


def test_xls_quando_xlrd_existe(bordas):
    xlrd = pytest.importorskip("xlrd")
    assert xlrd is not None
    # não há escritor de .xls no ambiente: o arquivo corrompido prova que o caminho do xlrd é tomado e que vira aviso
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600, "velha.xls",
               "application/vnd.ms-excel")
    r = _rodar(EMPRESA_A, [a])
    assert r.texto == "" and len(r.avisos) == 1 and "velha.xls" in r.avisos[0], r.avisos


# =====================================================================================================================
# (b) o path de OUTRA corretora vira aviso e NÃO é baixado
# =====================================================================================================================
def test_b_path_de_outra_corretora_nao_e_baixado(bordas):
    st = bordas["storage"]
    alheio = _anexo(st, EMPRESA_B, "chat-docs", "csv", _csv_ponto_e_virgula(), "da-outra.csv", "text/csv")
    r = _rodar(EMPRESA_A, [alheio])
    assert st.baixados == [], f"baixou o arquivo de outra corretora: {st.baixados}"
    assert r.texto == "" and r.avisos == ["O anexo «da-outra.csv» não pertence a esta corretora e não foi lido."]
    assert r.lidos[0]["lido"] is False


def test_b_fuga_por_ponto_ponto_e_bucket_alheio_tambem(bordas):
    st = bordas["storage"]
    fuga = {"bucket": "chat-docs", "path": f"{EMPRESA_A}/../{EMPRESA_B}/x/y.csv", "nome": "y.csv", "mime": "text/csv"}
    balde = {**_anexo(st, EMPRESA_A, "portal-evidence", "pdf", _pdf(), "evidencia.pdf", "application/pdf")}
    sem_empresa = _anexo(st, EMPRESA_A, "chat-docs", "csv", _csv_ponto_e_virgula(), "a.csv", "text/csv", n=7)
    r = _rodar(EMPRESA_A, [fuga, balde])
    r2 = _rodar("", [sem_empresa])
    assert st.baixados == [] and len(r.avisos) == 2 and len(r2.avisos) == 1, (st.baixados, r.avisos, r2.avisos)


def test_b_controle_sem_a_conferencia_o_path_alheio_seria_baixado(bordas, monkeypatch):
    monkeypatch.setattr(CA, "path_da_corretora", lambda path, company_id: True)
    st = bordas["storage"]
    alheio = _anexo(st, EMPRESA_B, "chat-docs", "csv", _csv_ponto_e_virgula(), "da-outra.csv", "text/csv")
    _rodar(EMPRESA_A, [alheio])
    assert st.baixados == [("chat-docs", alheio["path"])]


# =====================================================================================================================
# (c) 21 MB vira aviso
# =====================================================================================================================
def test_c_arquivo_de_21_mb_vira_aviso(bordas):
    st = bordas["storage"]
    grande = b"%PDF-1.4\n" + b"0" * (21 * 1024 * 1024)
    declarado = _anexo(st, EMPRESA_A, "chat-docs", "pdf", grande, "enorme.pdf", "application/pdf", n=1)
    mentiroso = {**_anexo(st, EMPRESA_A, "chat-docs", "pdf", grande, "mentiu.pdf", "application/pdf", n=2), "bytes": 10}
    r = _rodar(EMPRESA_A, [declarado, mentiroso])
    assert r.avisos == ["O arquivo «enorme.pdf» passa de 20 MB e não foi lido.",
                        "O arquivo «mentiu.pdf» passa de 20 MB e não foi lido."], r.avisos
    assert st.baixados == [("chat-docs", mentiroso["path"])], "o declarado > 20 MB nem é baixado"
    assert r.texto == ""


# =====================================================================================================================
# (d) a imagem NUNCA vira URL externa — só data URL
# =====================================================================================================================
def test_d_imagem_vira_data_url_jpeg_e_a_descricao_mascarada(bordas):
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-media", "png", _png(3000), "para-choque.png", "image/png")
    r = _rodar(EMPRESA_A, [a])
    assert r.avisos == [] and len(r.imagens) == 1, r.avisos
    assert r.imagens[0].startswith("data:image/jpeg;base64,"), r.imagens[0][:40]
    import base64

    from PIL import Image

    jpeg = Image.open(io.BytesIO(base64.b64decode(r.imagens[0].split(",", 1)[1])))
    assert max(jpeg.size) <= CA.LADO_MAXIMO_DA_IMAGEM and jpeg.format == "JPEG", jpeg.size
    enviados = bordas["llm"].recebido
    assert len(enviados) == 1
    urls = [p["image_url"]["url"] for m in enviados[0] for p in (m.content if isinstance(m.content, list) else [])
            if isinstance(p, dict) and p.get("type") == "image_url"]
    assert urls == [r.imagens[0]], "o provedor recebe EXATAMENTE o data URL"
    assert all(u.startswith("data:image/") for u in urls) and not any("/api/storage" in u or "http" in u for u in urls)
    assert r.descricoes and CPF not in r.descricoes[0] and "[CPF]" in r.descricoes[0], r.descricoes
    bloco = r.bloco_do_turno()
    assert VS.imagem_ja_descrita(bloco) and "para-choque.png" in bloco and CPF not in bloco


def test_d_describe_image_recusa_caminho_relativo_nosso(bordas):
    """O caminho antigo (`/api/storage/...`) não vai mais ao provedor: é recusado ANTES da chamada paga."""
    out = asyncio.run(VS.describe_image("/api/storage/chat-media/x/y.png", company_id=EMPRESA_A))
    assert out is None and bordas["llm"].recebido == []
    out2 = asyncio.run(VS.describe_image(company_id=EMPRESA_A, image_bytes=_png(), mime="image/png"))
    assert out2 and bordas["llm"].recebido[-1][1].content[1]["image_url"]["url"].startswith("data:image/png;base64,")


def test_heic_vira_jpeg_quando_pillow_heif_existe(bordas):
    pytest.importorskip("pillow_heif")
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-media", "heic", _heic(), "IMG_0001.HEIC", "image/heic")
    r = _rodar(EMPRESA_A, [a])
    assert r.avisos == [] and r.imagens and r.imagens[0].startswith("data:image/jpeg;base64,"), r.avisos


def test_heic_sem_pillow_heif_vira_aviso_humano(bordas, monkeypatch):
    import builtins

    real = builtins.__import__

    def _sem_heif(nome, *a, **k):
        if nome == "pillow_heif":
            raise ImportError("sem pillow_heif (teste)")
        return real(nome, *a, **k)

    monkeypatch.setattr(builtins, "__import__", _sem_heif)
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-media", "heic", b"\0\0\0\x18ftypheic" + b"\0" * 64, "IMG_9.heic", "image/heic")
    r = _rodar(EMPRESA_A, [a])
    assert r.imagens == [] and r.avisos and "HEIC" in r.avisos[0] and "JPG" in r.avisos[0], r.avisos


# =====================================================================================================================
# (e) a falha de um arquivo não derruba os outros — e os outros formatos
# =====================================================================================================================
def test_e_um_arquivo_quebrado_nao_derruba_os_outros(bordas, monkeypatch):
    st = bordas["storage"]
    quebrado = _anexo(st, EMPRESA_A, "chat-docs", "xlsx", b"PK\x03\x04 isto nao e um xlsx", "quebrado.xlsx",
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", n=1)
    sumiu = {"bucket": "chat-docs", "path": _path(EMPRESA_A, "chat-docs", "pdf", 9), "nome": "sumiu.pdf",
             "mime": "application/pdf", "bytes": 100}
    bom = _anexo(st, EMPRESA_A, "chat-docs", "csv", _csv_ponto_e_virgula(), "bom.csv", "text/csv", n=2)
    foto = _anexo(st, EMPRESA_A, "chat-media", "png", _png(), "foto.png", "image/png", n=3)

    async def _visao_cai(*_a, **_k):
        raise RuntimeError("provedor fora do ar")

    monkeypatch.setattr(CA, "_descrever_imagem", _visao_cai)
    linhas = []

    async def _progresso(frase):
        linhas.append(frase)
        if "bom.csv" in frase:
            raise RuntimeError("a tela caiu")

    r = _rodar(EMPRESA_A, [quebrado, sumiu, bom, foto], ao_progredir=_progresso)
    assert "[DOCUMENTO ANEXADO: bom.csv]" in r.texto and r.imagens, (r.texto[:200], r.avisos)
    assert any("quebrado.xlsx" in a for a in r.avisos) and any("sumiu.pdf" in a for a in r.avisos), r.avisos
    assert len(r.avisos) == 2 and r.descricoes == [], r.avisos
    assert [x["lido"] for x in r.lidos] == [False, False, True, True]
    assert sorted(linhas) == sorted(["Lendo quebrado.xlsx…", "Lendo sumiu.pdf…", "Lendo bom.csv…",
                                     "Olhando a foto foto.png…"]), linhas


def test_e_orcamento_estourado_vira_aviso_sem_travar_o_turno(bordas, monkeypatch):
    monkeypatch.setattr(CA, "ORCAMENTO_TOTAL_S", 0.3)
    st = bordas["storage"]
    lento = _anexo(st, EMPRESA_A, "chat-docs", "txt", b"texto", "lento.txt", "text/plain", n=1)
    rapido = _anexo(st, EMPRESA_A, "chat-docs", "txt", b"ola mundo", "rapido.txt", "text/plain", n=2)
    baixar = st.baixar

    async def _devagar(bucket, path):
        if path == lento["path"]:
            await asyncio.sleep(5)
        return await baixar(bucket, path)

    monkeypatch.setattr(CA, "_baixar_do_storage", _devagar)
    r = _rodar(EMPRESA_A, [lento, rapido])
    assert "ola mundo" in r.texto and any("Demorei demais" in a and "lento.txt" in a for a in r.avisos), r.avisos


def test_pdf_pelo_leitor_da_130b_mascarado(bordas):
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "pdf", _pdf(), "ata.pdf", "application/pdf")
    r = _rodar(EMPRESA_A, [a])
    assert r.avisos == [], r.avisos
    assert r.texto.startswith("[DOCUMENTO ANEXADO: ata.pdf]") and "fachada" in r.texto, r.texto[:300]
    assert _sem_pii(r.texto) == [], _sem_pii(r.texto)


def test_docx_e_pptx_por_bytes_mascarados(bordas):
    st = bordas["storage"]
    d = _anexo(st, EMPRESA_A, "chat-docs", "docx", _docx(), "proposta.docx",
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", n=1)
    p = _anexo(st, EMPRESA_A, "chat-docs", "pptx", _pptx(), "apresentacao.pptx",
               "application/vnd.openxmlformats-officedocument.presentationml.presentation", n=2)
    r = _rodar(EMPRESA_A, [d, p])
    assert r.avisos == [], r.avisos
    assert "renovação do seguro auto" in r.texto and "Carteira 2026" in r.texto and "[slide 1]" in r.texto, r.texto
    assert _sem_pii(r.texto) == [], _sem_pii(r.texto)


def test_audio_por_bytes_transcrito_e_mascarado(bordas):
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "opus", b"OggS" + b"\0" * 200, "nota de voz.opus", "audio/ogg")
    r = _rodar(EMPRESA_A, [a])
    assert r.avisos == [] and len(r.transcricoes) == 1, r.avisos
    assert "sinistro do carro" in r.transcricoes[0] and _sem_pii(r.transcricoes[0]) == [], r.transcricoes
    nome_tmp, modelo = bordas["api"].recebidos[0]
    assert nome_tmp.endswith(".ogg") and modelo == "gpt-transcribe", "o .opus vai à API como .ogg"
    assert "[ÁUDIO ANEXADO: nota de voz.opus — transcrição]" in r.bloco_do_turno()


def test_nome_do_arquivo_com_cpf_sai_mascarado(bordas):
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "txt", b"conteudo qualquer", f"rg {CPF}.txt", "text/plain")
    r = _rodar(EMPRESA_A, [a])
    tudo = r.bloco_do_turno() + repr(r.lidos)
    assert CPF not in tudo and "conteudo qualquer" in r.texto, tudo[:300]


def test_mais_de_dez_anexos_le_os_dez_primeiros(bordas):
    st = bordas["storage"]
    anexos = [_anexo(st, EMPRESA_A, "chat-docs", "txt", f"arquivo {i}".encode(), f"a{i}.txt", "text/plain", n=i)
              for i in range(12)]
    r = _rodar(EMPRESA_A, anexos)
    assert len(r.lidos) == 10 and r.avisos == ["Li só os 10 primeiros arquivos desta mensagem."], r.avisos


def test_lista_vazia_e_lixo_nao_lancam(bordas):
    assert _rodar(EMPRESA_A, []).bloco_do_turno() == ""
    r = _rodar(EMPRESA_A, [None, "x", 3])  # type: ignore[list-item]
    assert r.lidos == [] and r.avisos == []


# =====================================================================================================================
# 🔴 O CONSERTO ÚNICO (14/10) — os achados do red team (B1, B2, P5, P6). Cada guarda tem a sua MUTAÇÃO, compilada
# numa cópia em memória (a árvore compartilhada nunca é escrita — lição da SPEC-121).
# =====================================================================================================================
import struct  # noqa: E402
import time as _t  # noqa: E402
import tracemalloc  # noqa: E402
import types  # noqa: E402

ANEXOS_PY = BACKEND / "app" / "services" / "chat_anexos.py"
VISION_PY = BACKEND / "app" / "services" / "vision_service.py"
#: 9 dígitos seguidos (com ou sem pontuação) recalculam um CPF inteiro: os 2 últimos são dígitos verificadores
_NOVE_DIGITOS = re.compile(r"(?:\d[.\-\s]?){9,}")


def _mutado(caminho: Path, nome: str, de: str, para: str):
    fonte = caminho.read_text(encoding="utf-8")
    assert fonte.count(de) == 1, f"a âncora da mutação sumiu de {caminho.name}: {de[:70]!r}"
    mod = types.ModuleType(nome)
    mod.__file__ = str(caminho)
    mod.__package__ = "app.services"
    sys.modules[nome] = mod
    try:
        exec(compile(fonte.replace(de, para), str(caminho), "exec"), mod.__dict__)
    finally:
        sys.modules.pop(nome, None)
    return mod


def _csv_com_obs_longa(posicao: int, cpf: str) -> bytes:
    obs = "x" * (posicao - 4) + "CPF " + cpf + " e o resto da observacao do corretor"
    return f"Segurado,Obs\nPessoa Ficticia,{obs}\n".encode("utf-8")


def _vazamentos_da_celula_longa(mod) -> list:
    """O CPF começando em cada posição 50..70 de uma célula longa — pelo leitor e pela máscara REAIS do módulo."""
    vaza = []
    for pos in range(50, 71):
        for cpf in (CPF, CPF_DIGITOS):
            saida = mod._mascarar(mod.tabela_da_planilha(_csv_com_obs_longa(pos, cpf), "csv"))
            m = _NOVE_DIGITOS.search(saida or "")
            if m:
                vaza.append((pos, cpf, m.group(0)))
    return vaza


def test_conserto_b1_a_celula_longa_e_mascarada_antes_do_corte(bordas):
    """📊 red team B1: o corte em 77 caracteres vinha ANTES da máscara — 4 de 7 posições vazavam 9–10 dígitos."""
    assert _vazamentos_da_celula_longa(CA) == []
    # pelo contrato inteiro (`preparar`), uma posição que vazava antes
    st = bordas["storage"]
    a = _anexo(st, EMPRESA_A, "chat-docs", "csv", _csv_com_obs_longa(64, CPF), "obs.csv", "text/csv")
    r = _rodar(EMPRESA_A, [a])
    assert r.avisos == [] and not _NOVE_DIGITOS.search(r.texto) and "Pessoa Ficticia" in r.texto, r.texto
    # 🔴 MUTAÇÃO — cortar antes de mascarar → VERMELHO
    mutado = _mutado(ANEXOS_PY, "app.services._anexos_corte_antes_132a",
                     "    if len(t) <= 80:\n        return t\n    mascarado = _mascarar(t)",
                     "    if len(t) <= 80:\n        return t\n    t = t[:77] + '...'\n    mascarado = _mascarar(t)")
    vaza = _vazamentos_da_celula_longa(mutado)
    print("📊 MUTADO (corte antes da máscara): %d vazamentos, ex.: %s" % (len(vaza), vaza[:2]))
    assert vaza, "a mutação passou — o guarda B1 é carimbo"


def _csv_largo() -> bytes:
    """📊 o CSV do red team: 1 linha com 10.000 vírgulas + 2.000 linhas curtas = 13,7 KB."""
    return ("a" + "," * 10000 + "b\n" + "x\n" * 2000).encode()


def _xlsx_largo(n: int = 1600) -> bytes:
    """📊 o xlsx do red team: por linha, 1 célula em A e 1 em XFD, dimensão A1:XFDn = 12,6 KB com 1.600 linhas."""
    linhas = "".join(f'<row r="{i}"><c r="A{i}" t="inlineStr"><is><t>a</t></is></c><c r="XFD{i}"><v>1</v></c></row>'
                     for i in range(1, n + 1))
    return _xlsx_cru(f'<dimension ref="A1:XFD{n}"/><sheetData>{linhas}</sheetData>')


def _xlsx_cru(miolo: str) -> bytes:
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rel = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
                   'content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
                   'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/'
                   'workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+'
                   'xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-'
                   'officedocument.spreadsheetml.worksheet+xml"/></Types>')
        z.writestr("_rels/.rels", f'<?xml version="1.0" encoding="UTF-8"?><Relationships {rel}><Relationship Id="rId1" '
                   'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                   'Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", f'<?xml version="1.0" encoding="UTF-8"?><workbook {ns} xmlns:r="http://schemas.'
                   'openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="S" sheetId="1" '
                   'r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", f'<?xml version="1.0" encoding="UTF-8"?><Relationships {rel}>'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                   'worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   f'<worksheet {ns}>{miolo}</worksheet>')
    return b.getvalue()


def _pico_mb(fn) -> float:
    tracemalloc.start()
    try:
        fn()
        return tracemalloc.get_traced_memory()[1] / 2 ** 20
    finally:
        tracemalloc.stop()


def test_conserto_b2_o_leitor_tabular_tem_teto_de_memoria():
    """📊 red team B2: 13,7 KB de CSV pediam 154 MB; 12,6 KB de xlsx pediam 407 MB (8 B × largura × linhas)."""
    csv_, xlsx = _csv_largo(), _xlsx_largo()
    pico_csv = _pico_mb(lambda: CA.tabela_da_planilha(csv_, "csv"))
    pico_xlsx = _pico_mb(lambda: CA.tabela_da_planilha(xlsx, "xlsx"))
    print("📊 pico: csv %.1f KB → %.1f MB · xlsx %.1f KB → %.1f MB" % (len(csv_) / 1024, pico_csv, len(xlsx) / 1024,
                                                                      pico_xlsx))
    assert pico_csv <= 30 and pico_xlsx <= 60, (pico_csv, pico_xlsx)
    t = CA.tabela_da_planilha(csv_, "csv")
    assert "2000 linhas de dados × 10001 colunas" in t and f"(li só as primeiras {CA.COLUNAS_LIDAS} colunas)" in t, \
        t[:400]
    assert f"(li só as primeiras {CA.COLUNAS_LIDAS} colunas)" in CA.tabela_da_planilha(xlsx, "xlsx")
    # 🔴 MUTAÇÃO — guardar cada linha COMPLETADA até a mais larga (o que o leitor de antes fazia) → VERMELHO
    mutado = _mutado(ANEXOS_PY, "app.services._anexos_guarda_tudo_132a",
                     "            amostra.append(celulas[:COLUNAS_DA_AMOSTRA])",
                     "            pass\n        amostra.append(list(ln) + [None] * (largura_real - len(ln)))")
    m_csv = _pico_mb(lambda: mutado.tabela_da_planilha(csv_, "csv"))
    print("📊 MUTADO (guarda as linhas inteiras): csv %.1f MB" % m_csv)
    assert m_csv > 30, "a mutação passou — o guarda B2 é carimbo"


def test_conserto_b2_teto_de_celulas_e_prazo_cooperativo():
    linhas = "Ramo,Valor\n" + "".join(f"Auto,{i}\n" for i in range(1, 120_001))
    t = CA.tabela_da_planilha(linhas.encode(), "csv")
    # 💭 200 mil células / 2 colunas: a leitura para em ~100 mil linhas e diz "+"
    assert re.search(r"— \d+\+ linhas de dados × 2 colunas", t) and "(a planilha é grande: li só o começo dela)" in t, \
        t[:300]
    with pytest.raises(CA.PrazoDaLeituraEstourado):
        CA.tabela_da_planilha(linhas.encode(), "csv", prazo=_t.monotonic() - 1)


def test_conserto_b2_xlsx_com_linha_forjada_e_zip_bomba_viram_aviso(bordas):
    st = bordas["storage"]
    forjada = _xlsx_cru('<sheetData><row r="1">' + "<c/>" * (CA.CELULAS_POR_LINHA_DO_XML + 5) + "</row></sheetData>")
    a = _anexo(st, EMPRESA_A, "chat-docs", "xlsx", forjada, "forjada.xlsx",
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", n=1)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", b"<w:t>a</w:t>" * (CA.DESCOMPRIMIDO_POR_PARTE_DO_DOCUMENTO // 12 + 10))
    d = _anexo(st, EMPRESA_A, "chat-docs", "docx", buf.getvalue(), "bomba.docx",
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", n=2)
    print("📊 xlsx forjado %.1f KB · docx bomba %.1f KB" % (len(forjada) / 1024, len(buf.getvalue()) / 1024))
    r = _rodar(EMPRESA_A, [a, d])
    assert any("forjada.xlsx" in x and "grande demais" in x for x in r.avisos), r.avisos
    assert any("bomba.docx" in x and "grande demais" in x for x in r.avisos), r.avisos
    with pytest.raises(CA.ArquivoGrandeDemais):
        CA.tabela_da_planilha(forjada, "xlsx")
    # 🔴 MUTAÇÃO — sem a conferência da linha, o xlsx forjado É lido (nenhuma recusa) → VERMELHO
    mutado = _mutado(ANEXOS_PY, "app.services._anexos_sem_linha_132a",
                     "                    if na_linha > CELULAS_POR_LINHA_DO_XML:\n                        raise",
                     "                    if False:\n                        raise")
    try:
        mutado.tabela_da_planilha(forjada, "xlsx")
        recusou = False
    except mutado.ArquivoGrandeDemais:
        recusou = True
    assert not recusou, "a mutação passou — o guarda da linha forjada é carimbo"


def _png_grande(lado: int = 7500) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("1", (lado, lado), 0).save(buf, format="PNG")
    return buf.getvalue()


def test_conserto_b2_foto_gigante_vira_aviso_sem_decodificar():
    """📊 red team B2: PNG 1-bit de 81 Mpx em 10 KB → +515 MB; 169 Mpx/20 KB abria só com um aviso do Pillow."""
    png = _png_grande()
    inicio = _t.monotonic()
    jpeg, _tipo, aviso = CA._imagem_para_jpeg(png, "image/png", "foto.png")
    dt = _t.monotonic() - inicio
    print("📊 PNG 7500² (%.0f Mpx) em %.0f KB → %s em %.3f s" % (7500 * 7500 / 1e6, len(png) / 1024,
                                                               "aviso" if aviso else "jpeg", dt))
    assert jpeg is None and aviso and aviso.startswith("Esta foto é grande demais"), aviso
    assert dt < 1.0, "a foto foi decodificada (%.2f s)" % dt
    # o JPEG grande é lido REDUZIDO (draft): sai com o lado ≤ 2048
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (4000, 3000), (10, 120, 200)).save(buf, format="JPEG", quality=80)
    j, _t2, av = CA._imagem_para_jpeg(buf.getvalue(), "image/jpeg", "grande.jpg")
    assert j and av is None and max(Image.open(io.BytesIO(j)).size) <= CA.LADO_MAXIMO_DA_IMAGEM
    # 🔴 MUTAÇÃO — sem a conferência dos pixels, a foto gigante é decodificada inteira → VERMELHO
    mutado = _mutado(ANEXOS_PY, "app.services._anexos_sem_pixels_132a",
                     "            if not e_jpeg and largura * altura > PIXELS_MAXIMOS:",
                     "            if False:")
    jm, _tm, avm = mutado._imagem_para_jpeg(png, "image/png", "foto.png")
    assert jm is not None and avm is None, "a mutação passou — o guarda da foto gigante é carimbo"


def test_conserto_b2_dez_anexos_ruins_nao_prendem_o_poco_do_processo(bordas, monkeypatch):
    """O poço do processo (`EXECUTOR_THREADS`) atende o WhatsApp de todas as corretoras: dez planilhas lentas no
    MESMO turno não podem segurá-lo. E o orçamento estourado CANCELA o que ainda estava na fila."""
    import concurrent.futures as cf

    st = bordas["storage"]
    anexos = [_anexo(st, EMPRESA_A, "chat-docs", "csv", b"a,b\n1,2\n", f"lenta{i}.csv", "text/csv", n=i)
              for i in range(10)]

    def _medir(mod) -> tuple:
        comecaram = []

        def _lenta(dados, tipo, *, prazo=None):  # a borda: um leitor que demora (o motor do poço é o real)
            comecaram.append(1)
            _t.sleep(1.5)
            return "x"

        monkeypatch.setattr(mod, "tabela_da_planilha", _lenta)
        monkeypatch.setattr(mod, "ORCAMENTO_TOTAL_S", 0.3)
        if mod is not CA:
            monkeypatch.setattr(mod, "_baixar_do_storage", st.baixar)

        async def _cenario():
            asyncio.get_running_loop().set_default_executor(cf.ThreadPoolExecutor(max_workers=2))
            leitura = asyncio.ensure_future(mod.preparar(EMPRESA_A, anexos))
            await asyncio.sleep(0.1)
            t0 = _t.monotonic()
            try:
                await asyncio.wait_for(asyncio.to_thread(lambda: "ok"), 0.8)
                sonda = _t.monotonic() - t0
            except asyncio.TimeoutError:
                sonda = None
            r = await leitura
            await asyncio.sleep(2.0)
            return sonda, r

        sonda, r = asyncio.run(_cenario())
        return sonda, len(comecaram), r

    sonda, comecaram, r = _medir(CA)
    print("📊 outro to_thread do processo esperou %s s · leituras que COMEÇARAM: %d de 10" % (sonda, comecaram))
    assert sonda is not None and sonda < 0.5, sonda
    assert comecaram == CA.TRABALHADORES_DOS_ANEXOS, "o orçamento estourado não cancelou a fila: %d" % comecaram
    assert sum("Demorei demais" in a for a in r.avisos) == 10, r.avisos
    # 🔴 MUTAÇÃO — as leituras no poço do PROCESSO (o `to_thread` de antes) → a sonda espera → VERMELHO
    mutado = _mutado(ANEXOS_PY, "app.services._anexos_poco_padrao_132a",
                     "    return await laco.run_in_executor(_poco_dos_anexos(), ",
                     "    return await laco.run_in_executor(None, ")
    sonda_m, _c, _r = _medir(mutado)
    print("📊 MUTADO (poço do processo): a sonda esperou %s" % sonda_m)
    assert sonda_m is None or sonda_m >= 0.5, "a mutação passou — o guarda do poço é carimbo"


def _m4a(segundos: int) -> bytes:
    """Um .m4a mínimo: `ftyp` + `moov/mvhd` (escala 1000) — a duração mora no cabeçalho."""
    mvhd_corpo = b"\0\0\0\0" + b"\0" * 8 + struct.pack(">II", 1000, segundos * 1000) + b"\0" * 80
    mvhd = struct.pack(">I", 8 + len(mvhd_corpo)) + b"mvhd" + mvhd_corpo
    moov = struct.pack(">I", 8 + len(mvhd)) + b"moov" + mvhd
    ftyp = struct.pack(">I", 16) + b"ftypM4A \0\0\0\0"
    return ftyp + moov + b"\0" * 64


def test_conserto_p6_audio_de_mais_de_15_min_vira_aviso_sem_transcrever(bordas, monkeypatch):
    st = bordas["storage"]
    longo = _anexo(st, EMPRESA_A, "chat-docs", "m4a", _m4a(16 * 60), "reuniao.m4a", "audio/mp4", n=1)
    curto = _anexo(st, EMPRESA_A, "chat-docs", "m4a", _m4a(14 * 60), "nota.m4a", "audio/mp4", n=2)
    assert CA.duracao_do_audio(_m4a(16 * 60)) == 960.0
    r = _rodar(EMPRESA_A, [longo, curto])
    assert any(a.startswith("Este áudio passa de 15 minutos («reuniao.m4a»)") for a in r.avisos), r.avisos
    assert len(bordas["api"].recebidos) == 1 and len(r.transcricoes) == 1, "só o áudio de 14 min foi à transcrição"
    # 🔴 MUTAÇÃO — sem o teto de duração, o áudio de 16 min vai à transcrição (e ao ledger) → VERMELHO
    mutado = _mutado(ANEXOS_PY, "app.services._anexos_sem_duracao_132a",
                     "        if duracao is not None and duracao > SEGUNDOS_MAXIMOS_DO_AUDIO:",
                     "        if False:")
    monkeypatch.setattr(mutado, "_baixar_do_storage", st.baixar)
    monkeypatch.setattr(mutado, "_servico_de_audio", CA._servico_de_audio)
    rm = asyncio.run(mutado.preparar(EMPRESA_A, [longo]))
    assert not any("15 minutos" in a for a in rm.avisos) and rm.transcricoes, \
        "a mutação passou — o guarda do áudio longo é carimbo"


def test_conserto_p5_o_download_antigo_so_vai_ao_storage_da_plataforma(monkeypatch):
    """red team P5 (SSRF): `fileUrl` do corpo virava `httpx.get` com `follow_redirects=True` para qualquer host."""
    import httpx

    from app.core import settings

    monkeypatch.setattr(settings, "SUPABASE_URL", "https://projeto-ficticio.supabase.co", raising=False)
    criados = []

    class _Cliente:
        def __init__(self, *a, **k):
            criados.append(k)

        async def __aenter__(self):
            return self

        async def __aexit__(self, *a):
            return False

        async def get(self, url):
            return httpx.Response(200, content=b"%PDF", request=httpx.Request("GET", url))

    monkeypatch.setattr(httpx, "AsyncClient", _Cliente)

    def _baixar(mod, url):
        return asyncio.run(mod.baixar_documento(url))

    for ruim in ("http://169.254.169.254/latest/meta-data/", "https://interno.local/x.pdf",
                 "http://projeto-ficticio.supabase.co/storage/v1/object/a.pdf",
                 "https://projeto-ficticio.supabase.co.atacante.example/a.pdf"):
        assert _baixar(VS, ruim) is None, ruim
    assert criados == [], "nenhum cliente HTTP para um host fora do Storage"
    assert _baixar(VS, "https://projeto-ficticio.supabase.co/storage/v1/object/sign/chat-docs/a.pdf?token=t") == b"%PDF"
    assert criados and criados[-1].get("follow_redirects") is False
    # 🔴 MUTAÇÃO — sem a lista do host → o endereço interno é baixado → VERMELHO
    mutado = _mutado(VISION_PY, "app.services._vision_sem_lista_132a",
                     "        if alvo.scheme != \"https\" or not alvo.hostname or alvo.hostname != nosso.hostname "
                     "or alvo.port != nosso.port:",
                     "        if False:")
    assert _baixar(mutado, "http://169.254.169.254/latest/meta-data/") == b"%PDF", \
        "a mutação passou — o guarda do SSRF é carimbo"
