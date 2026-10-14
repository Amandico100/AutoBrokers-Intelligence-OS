"""SPEC-132-A · A2 — OS ANEXOS DO CHAT PRINCIPAL, LIDOS POR BYTES.

O fio (SPEC-132-A §2):
    composer → POST /api/upload → {bucket, path, nome, mime, bytes}
    → POST /api/chat/stream {attachments[]} → BFF (path com o prefixo DESTA corretora, senão 403)
    → backend /chat/stream → _gerar → **chat_anexos.preparar(company_id, attachments)** → bloco do turno → modelo

📊 B3/B4 (14/10/2026): o anexo "dava erro" por dois motivos encadeados — a tela lia `publicUrl` (que o upload nunca
devolveu) e, quando havia URL, era `/api/storage/...` RELATIVA e atrás da sessão: o `httpx.get` do backend e o
provedor de visão nunca a alcançavam. Aqui o backend baixa os BYTES do Supabase Storage com a service role
(D-132A-05) e o provedor de IA recebe só data URL — nunca URL nossa.

Cada formato pelo leitor que JÁ existe (CLAUDE.md §5 — nenhum leitor paralelo):
    imagem (jpeg/png/gif/webp/HEIC) → JPEG ≤ 2048 px em data URL + a descrição do papel `visao` (`describe_image`)
    PDF                              → `vision_service.conteudo_do_anexo` (o leitor da SPEC-130-B, audiência corretor)
    xlsx/xls/csv                     → o leitor TABULAR deste módulo (abas, cabeçalho, amostra, totais)
    docx/pptx                        → docling quando instalado; senão python-docx / o XML do pacote Office
    txt/md                           → o texto
    áudio                            → `AudioService.transcribe_audio_bytes` (papel `transcricao`)
🔴 TUDO que é texto sai pela máscara da 130-B (`vision_service.mascarar`, fail-closed): sem a máscara, o texto não sai.
🔴 Defesa em profundidade: o BFF já conferiu o prefixo; aqui ele é conferido DE NOVO contra o `company_id` do turno.
⛔ Nunca lança. O que falha vira AVISO em português. Nunca loga conteúdo, nome de arquivo nem CPF — só contagens.
"""
from __future__ import annotations

import asyncio
import codecs
import csv
import functools
import io
import logging
import re
import struct
import threading
import time as _relogio
import unicodedata
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Awaitable, Callable, Dict, Iterable, List, Optional, Tuple

logger = logging.getLogger(__name__)

#: Cada arquivo cabe em 20 MB (o mesmo teto do upload — `app/api/upload/route.ts`).
LIMITE_DE_BYTES = 20 * 1024 * 1024
#: Até 10 arquivos por mensagem (o mesmo teto do BFF).
MAX_ANEXOS = 10
#: 💭 Orçamento TOTAL da leitura de todos os anexos de um turno (SPEC-132-A §5).
ORCAMENTO_TOTAL_S = 45.0
#: Os buckets do chat. `avatars`, `voice-messages` e `portal-evidence` nunca são lidos por aqui.
BUCKETS_DO_CHAT = frozenset({"chat-media", "chat-docs"})
#: 💭 A amostra da planilha: até 50 linhas por aba.
LINHAS_DA_AMOSTRA = 50
#: 💭 Colunas mostradas na amostra (o resto é contado, não mostrado).
COLUNAS_DA_AMOSTRA = 30
#: 💭 Abas lidas por planilha e linhas lidas por aba (para contar e somar).
ABAS_LIDAS = 10
LINHAS_LIDAS_POR_ABA = 50_000
#: O lado maior da imagem que vai ao modelo.
LADO_MAXIMO_DA_IMAGEM = 2048
#: 💭 O teto do texto de TODOS os anexos juntos (cada documento já vem truncado em 20.000 pela 130-B).
TETO_DO_TEXTO = 60_000

# 🔴 Conserto da SPEC-132-A (red team B2): os tetos de RECURSO, não só de bytes do arquivo.
# 📊 14/10 (`tabela_da_planilha` real + tracemalloc): um CSV de 14 KB com uma linha de 10.000 vírgulas pedia 154 MB, e um
# xlsx de 12,6 KB com uma célula em XFD pedia 407 MB — o leitor completava TODA linha até a mais larga e guardava tudo.
#: 💭 Colunas lidas por aba (as de depois são CONTADAS, nunca guardadas).
COLUNAS_LIDAS = 200
#: 💭 Células lidas por arquivo (linhas × colunas, todas as abas juntas). Passou → a leitura para e diz "+".
CELULAS_POR_ARQUIVO = 200_000
#: 💭 Uma linha física de CSV maior que isto é cortada ANTES do parser (só as primeiras colunas são lidas).
CARACTERES_POR_LINHA_DO_CSV = 65_536
#: 💭 O pacote Office (xlsx/docx/pptx) DESCOMPRIMIDO: por parte, no total, e a razão que denuncia uma bomba de zip.
DESCOMPRIMIDO_POR_PARTE_DA_PLANILHA = 80 * 1024 * 1024
DESCOMPRIMIDO_POR_PARTE_DO_DOCUMENTO = 30 * 1024 * 1024
DESCOMPRIMIDO_DAS_STRINGS = 40 * 1024 * 1024
DESCOMPRIMIDO_TOTAL = 200 * 1024 * 1024
RAZAO_SUSPEITA = 100
#: Uma linha do Excel tem no máximo 16.384 colunas; mais do que 💭 20.000 `<c>` numa linha só = arquivo forjado.
CELULAS_POR_LINHA_DO_XML = 20_000
#: 💭 Pixels de uma foto que não é JPEG (PNG/GIF/WebP/HEIC decodificam INTEIROS). O JPEG acima disto é lido reduzido
#: (`draft`, sem decodificar o original inteiro) até o teto do próprio Pillow.
PIXELS_MAXIMOS = 50_000_000
#: 💭 Duração máxima de um áudio anexado (quando o cabeçalho a diz). O tamanho já tem o teto de 20 MB de todo arquivo.
SEGUNDOS_MAXIMOS_DO_AUDIO = 15 * 60
#: 💭 As leituras pesadas (Pillow, openpyxl, xlrd, python-docx) rodam num poço PRÓPRIO e pequeno: dez anexos ruins
#: nunca ocupam o poço do processo — o mesmo que atende o WhatsApp dos segurados (`EXECUTOR_THREADS`).
TRABALHADORES_DOS_ANEXOS = 3

AVISO_FOTO_GRANDE = ("Esta foto é grande demais para eu abrir («{nome}»). Pode mandar uma versão menor — um print da "
                     "tela serve?")
AVISO_AUDIO_LONGO = ("Este áudio passa de 15 minutos («{nome}») e não foi ouvido. Pode mandar um trecho menor, ou "
                     "escrever em poucas palavras o que ele diz?")


class ArquivoGrandeDemais(Exception):
    """O arquivo passa de um teto de recurso (descomprimido, linha forjada). Vira aviso — nunca derruba o turno."""


class PrazoDaLeituraEstourado(Exception):
    """O orçamento do turno acabou: a leitura PARA (cooperativa), e a thread volta ao poço."""


_POCO_DOS_ANEXOS: Optional[ThreadPoolExecutor] = None
_TRAVA_DO_POCO = threading.Lock()


def _poco_dos_anexos() -> ThreadPoolExecutor:
    global _POCO_DOS_ANEXOS
    if _POCO_DOS_ANEXOS is None:
        with _TRAVA_DO_POCO:
            if _POCO_DOS_ANEXOS is None:
                _POCO_DOS_ANEXOS = ThreadPoolExecutor(max_workers=TRABALHADORES_DOS_ANEXOS,
                                                      thread_name_prefix="anexos-do-chat")
    return _POCO_DOS_ANEXOS


async def _no_poco(fn: Callable[..., Any], *args: Any, **kwargs: Any) -> Any:
    """Roda a leitura pesada no poço DOS ANEXOS. Cancelada a espera (orçamento estourado), a que ainda estava na fila
    nem começa (o `Future` é cancelado junto); a que já corre para no próximo ponto de conferência do `prazo`."""
    laco = asyncio.get_running_loop()
    return await laco.run_in_executor(_poco_dos_anexos(), functools.partial(fn, *args, **kwargs))

_SEGMENTO_SEGURO = re.compile(r"^[A-Za-z0-9._-]{1,200}$")


@dataclass
class ResultadoDosAnexos:
    """O que os anexos de UM turno viram para o modelo.

    texto         blocos `[DOCUMENTO ANEXADO: …]…[FIM DO DOCUMENTO]` (PDF, planilha, Word, PowerPoint, texto) — MASCARADOS
    imagens       data URLs `data:image/jpeg;base64,…` (para um modelo que vê) — nunca URL nossa
    descricoes    a descrição de cada foto pelo papel `visao`, MASCARADA (para um modelo que não vê)
    transcricoes  o texto de cada áudio, MASCARADO
    avisos        o que não deu para ler, em português, para o corretor (e para o modelo dizer com naturalidade)
    lidos         um registro por anexo, SEM conteúdo: {nome (mascarado), mime, bytes, path, bucket, tipo, lido}
    """

    texto: str = ""
    imagens: List[str] = field(default_factory=list)
    descricoes: List[str] = field(default_factory=list)
    transcricoes: List[str] = field(default_factory=list)
    avisos: List[str] = field(default_factory=list)
    lidos: List[Dict[str, Any]] = field(default_factory=list)
    _nomes_das_descricoes: List[str] = field(default_factory=list, repr=False)
    _nomes_das_transcricoes: List[str] = field(default_factory=list, repr=False)

    def bloco_do_turno(self) -> str:
        """Tudo junto, pronto para entrar no texto do turno (o que a A3 anexa à mensagem do corretor).

        As descrições levam a MARCA do contexto visual (`[CONTEXTO VISUAL`), que o grafo reconhece para não descrever
        a foto de novo (`vision_service.imagem_ja_descrita`)."""
        partes: List[str] = []
        if self.texto:
            partes.append(self.texto)
        for nome, desc in zip(self._nomes_das_descricoes, self.descricoes):
            partes.append(f"[CONTEXTO VISUAL — imagem enviada pelo usuário: {nome}]:\n{desc}")
        for nome, transc in zip(self._nomes_das_transcricoes, self.transcricoes):
            partes.append(f"[ÁUDIO ANEXADO: {nome} — transcrição]\n{transc}\n[FIM DO ÁUDIO]")
        for aviso in self.avisos:
            partes.append(f"[AVISO INTERNO: {aviso} — diga isso ao usuário com naturalidade.]")
        return "\n\n".join(partes)


# =====================================================================================================================
# as bordas (os testes trocam SÓ estas: o Storage, a visão e o áudio)
# =====================================================================================================================
async def _baixar_do_storage(bucket: str, path: str) -> Optional[bytes]:
    """Os BYTES do objeto no Supabase Storage, com a service role do backend (o cliente que já existe)."""
    from app.core.database import get_supabase_client

    def _baixar() -> Optional[bytes]:
        dado = get_supabase_client().client.storage.from_(bucket).download(path)
        return bytes(dado) if dado else None

    return await asyncio.to_thread(_baixar)


async def _descrever_imagem(data_url: str, company_id: str) -> Optional[str]:
    """A descrição da foto pelo papel `visao` — o `describe_image` de sempre, agora com data URL."""
    from app.services import vision_service

    return await vision_service.describe_image(data_url, company_id=company_id, purpose_hint="Assistente")


def _servico_de_audio():
    """O `AudioService` do backend (papel `transcricao`)."""
    from app.core import settings
    from app.services.audio_service import AudioService

    return AudioService(openai_api_key=settings.OPENAI_API_KEY)


async def _transcrever_audio(dados: bytes, mime: str, nome: str, company_id: str) -> str:
    return await _servico_de_audio().transcribe_audio_bytes(dados, mime, nome=nome, company_id=company_id)


# =====================================================================================================================
# a máscara (a da 130-B; fail-closed)
# =====================================================================================================================
def _mascarar(texto: str) -> Optional[str]:
    from app.services.vision_service import mascarar

    return mascarar(texto)


def _nome_para_o_modelo(nome: str) -> str:
    try:
        from app.services.vision_service import nome_do_anexo_para_o_modelo

        return nome_do_anexo_para_o_modelo(nome, padrao="arquivo")
    except Exception:  # noqa: BLE001 — sem a máscara, nunca o nome cru
        return "arquivo"


def _truncar(texto: str) -> str:
    from app.services.vision_service import truncate_document_text

    return truncate_document_text(texto)


# =====================================================================================================================
# o tipo do anexo
# =====================================================================================================================
_EXT_IMAGEM = {"jpg", "jpeg", "png", "gif", "webp", "heic", "heif"}
_EXT_AUDIO = {"m4a", "mp3", "mpeg", "ogg", "oga", "opus", "wav", "webm", "aac"}
_MIME_PARA_TIPO = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xls",
    "text/csv": "csv",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "text/plain": "texto",
    "text/markdown": "texto",
}
_EXT_PARA_TIPO = {"pdf": "pdf", "xlsx": "xlsx", "xls": "xls", "csv": "csv", "docx": "docx", "pptx": "pptx",
                  "txt": "texto", "md": "texto", "markdown": "texto"}


def _extensao(nome: str) -> str:
    nome = str(nome or "")
    return nome.rsplit(".", 1)[-1].lower() if "." in nome else ""


def tipo_do_anexo(mime: str, nome: str = "", path: str = "") -> str:
    """imagem · pdf · xlsx · xls · csv · docx · pptx · texto · audio · '' (desconhecido). A extensão do NOME vence a do path."""
    ext = _extensao(nome) or _extensao(path)
    # o .csv do Windows chega como application/vnd.ms-excel: a extensão decide antes do MIME
    if ext in _EXT_PARA_TIPO:
        return _EXT_PARA_TIPO[ext]
    if ext in _EXT_IMAGEM:
        return "imagem"
    if ext in _EXT_AUDIO:
        return "audio"
    m = str(mime or "").split(";")[0].strip().lower()
    if m.startswith("image/"):
        return "imagem"
    if m.startswith("audio/"):
        return "audio"
    return _MIME_PARA_TIPO.get(m, "")


def path_da_corretora(path: Any, company_id: str) -> bool:
    """O objeto é DESTA corretora? Segmentos limpos, sem `..`, e o 1º é o `company_id` do turno."""
    if not isinstance(path, str) or not path or len(path) > 512 or not company_id:
        return False
    segmentos = path.split("/")
    if len(segmentos) < 2:
        return False
    if not all(_SEGMENTO_SEGURO.match(s) and s not in (".", "..") for s in segmentos):
        return False
    return segmentos[0] == company_id


# =====================================================================================================================
# imagem
# =====================================================================================================================
def _e_heic(dados: bytes, mime: str, nome: str) -> bool:
    m = str(mime or "").lower()
    return ("heic" in m or "heif" in m or _extensao(nome) in ("heic", "heif")
            or (dados[4:8] == b"ftyp" and dados[8:12] in (b"heic", b"heix", b"hevc", b"hevx", b"heim", b"heis",
                                                           b"mif1", b"msf1")))


AVISO_HEIC = ("Não consegui abrir a foto «{nome}» (formato HEIC, do iPhone). Pode mandar em JPG — ou tirar um print "
              "da tela e mandar o print?")


def _imagem_para_jpeg(dados: bytes, mime: str, nome: str) -> Tuple[Optional[bytes], str, Optional[str]]:
    """(bytes, mime, aviso). HEIC → JPEG; qualquer foto → lado ≤ 2048 px, orientação do EXIF, JPEG 85."""
    heic = _e_heic(dados, mime, nome)
    try:
        from PIL import Image, ImageOps
    except Exception:  # noqa: BLE001 — sem Pillow: só os formatos que o provedor lê direto
        if heic:
            return None, "", AVISO_HEIC
        tipo = str(mime or "image/jpeg").split(";")[0].strip().lower()
        if tipo in ("image/jpeg", "image/png", "image/gif", "image/webp") and len(dados) <= 5 * 1024 * 1024:
            return dados, tipo, None
        return None, "", "Não consegui preparar a foto «{nome}» para leitura. Pode mandar em tamanho menor?"
    if heic:
        try:
            import pillow_heif

            pillow_heif.register_heif_opener()
        except Exception:  # noqa: BLE001
            return None, "", AVISO_HEIC
    try:
        with Image.open(io.BytesIO(dados)) as original:
            # 🔴 red team B2: o tamanho vem do CABEÇALHO (o `open` não decodifica). 📊 14/10: um PNG 1-bit de 169 Mpx
            # e 20 KB abria com só um aviso do Pillow e pedia ~1 GB ao ser convertido.
            largura, altura = original.size
            e_jpeg = str(original.format or "").upper() in ("JPEG", "MPO")
            if not e_jpeg and largura * altura > PIXELS_MAXIMOS:
                return None, "", AVISO_FOTO_GRANDE
            original.seek(0)
            if e_jpeg:
                # o JPEG se decodifica JÁ REDUZIDO (1/2, 1/4, 1/8): a foto de 48 Mpx não vira 144 MB de pixels
                original.draft("RGB", (LADO_MAXIMO_DA_IMAGEM, LADO_MAXIMO_DA_IMAGEM))
            img = original
            # reduz ANTES de girar/converter: o `exif_transpose` e o RGBA decodificavam o original inteiro
            img.thumbnail((LADO_MAXIMO_DA_IMAGEM, LADO_MAXIMO_DA_IMAGEM))
            img = ImageOps.exif_transpose(img)
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
                fundo = Image.new("RGB", img.size, (255, 255, 255))
                fundo.paste(img, mask=img.split()[-1])
                img = fundo
            elif img.mode != "RGB":
                img = img.convert("RGB")
            saida = io.BytesIO()
            img.save(saida, format="JPEG", quality=85, optimize=True)
            return saida.getvalue(), "image/jpeg", None
    except Image.DecompressionBombError:
        return None, "", AVISO_FOTO_GRANDE
    except Exception as e:  # noqa: BLE001
        logger.warning("[ANEXOS] imagem não abriu (%s)", type(e).__name__)
        if heic:
            return None, "", AVISO_HEIC
        return None, "", "Não consegui abrir a foto «{nome}». Ela pode estar corrompida — tente mandar de novo."


# =====================================================================================================================
# áudio — a duração pelo CABEÇALHO (sem decodificar), para o teto de 💭 15 min
# =====================================================================================================================
def _duracao_wav(d: bytes) -> Optional[float]:
    if d[:4] != b"RIFF" or d[8:12] != b"WAVE":
        return None
    i, taxa, tamanho = 12, None, None
    while i + 8 <= len(d) and (taxa is None or tamanho is None):
        bloco, n = d[i:i + 4], struct.unpack("<I", d[i + 4:i + 8])[0]
        if bloco == b"fmt " and i + 20 <= len(d):
            taxa = struct.unpack("<I", d[i + 16:i + 20])[0]
        elif bloco == b"data":
            tamanho = min(n, len(d) - i - 8)
        i += 8 + n + (n & 1)
    return tamanho / taxa if taxa and tamanho is not None else None


def _duracao_mp4(d: bytes) -> Optional[float]:
    def _caixas(ini: int, fim: int):
        i = ini
        while i + 8 <= fim:
            n, tipo = struct.unpack(">I", d[i:i + 4])[0], d[i + 4:i + 8]
            cab = 8
            if n == 1 and i + 16 <= fim:
                n, cab = struct.unpack(">Q", d[i + 8:i + 16])[0], 16
            elif n == 0:
                n = fim - i
            if n < cab:
                return
            yield tipo, i + cab, min(i + n, fim)
            i += n

    if d[4:8] != b"ftyp":
        return None
    for tipo, ini, fim in _caixas(0, len(d)):
        if tipo != b"moov":
            continue
        for sub, s_ini, _s_fim in _caixas(ini, fim):
            if sub != b"mvhd":
                continue
            versao = d[s_ini]
            if versao == 1:
                escala, duracao = struct.unpack(">IQ", d[s_ini + 20:s_ini + 32])
            else:
                escala, duracao = struct.unpack(">II", d[s_ini + 12:s_ini + 20])
            return duracao / escala if escala else None
    return None


def _duracao_ogg(d: bytes) -> Optional[float]:
    if d[:4] != b"OggS":
        return None
    inicio = d[:4096]
    if b"OpusHead" in inicio:
        k = inicio.index(b"OpusHead")
        taxa, pre = 48000, struct.unpack("<H", inicio[k + 10:k + 12])[0]
    elif b"\x01vorbis" in inicio:
        k = inicio.index(b"\x01vorbis")
        taxa, pre = struct.unpack("<I", inicio[k + 12:k + 16])[0], 0
    else:
        return None
    ultima = d.rfind(b"OggS", max(0, len(d) - 65536))
    if ultima < 0 or ultima + 14 > len(d) or not taxa:
        return None
    granulo = struct.unpack("<q", d[ultima + 6:ultima + 14])[0]
    return max(0, granulo - pre) / taxa if granulo > 0 else None


def _vint(d: bytes, i: int, manter_marca: bool) -> Tuple[Optional[int], int]:
    if i >= len(d) or d[i] == 0:
        return None, i + 1
    n = 8 - d[i].bit_length() + 1
    v = d[i] if manter_marca else d[i] & (0xFF >> n)
    for b in d[i + 1:i + n]:
        v = (v << 8) | b
    return v, i + n


def _duracao_webm(d: bytes) -> Optional[float]:
    if d[:4] != b"\x1a\x45\xdf\xa3":
        return None
    k = d.find(b"\x15\x49\xa9\x66", 0, 1 << 20)  # Segment › Info
    if k < 0:
        return None
    tamanho, i = _vint(d, k + 4, False)
    fim = min(len(d), i + (tamanho or 0))
    escala, duracao = 1_000_000, None
    while i < fim:
        ident, i = _vint(d, i, True)
        n, i = _vint(d, i, False)
        if ident is None or n is None:
            return None
        if ident == 0x2AD7B1:
            escala = int.from_bytes(d[i:i + n], "big") or escala
        elif ident == 0x4489 and n in (4, 8):
            duracao = struct.unpack(">f" if n == 4 else ">d", d[i:i + n])[0]
        i += n
    return duracao * escala / 1e9 if duracao else None


_BITRATES_MP3 = {1: [0, 32, 40, 48, 56, 64, 80, 96, 112, 128, 160, 192, 224, 256, 320],
                 2: [0, 8, 16, 24, 32, 40, 48, 56, 64, 80, 96, 112, 128, 144, 160]}
_TAXAS_MP3 = {3: [44100, 48000, 32000], 2: [22050, 24000, 16000], 0: [11025, 12000, 8000]}


def _duracao_mp3(d: bytes) -> Optional[float]:
    i = 0
    if d[:3] == b"ID3" and len(d) >= 10:
        i = 10 + ((d[6] & 0x7F) << 21 | (d[7] & 0x7F) << 14 | (d[8] & 0x7F) << 7 | (d[9] & 0x7F))
    limite = min(len(d) - 4, i + 65536)
    while i < limite and not (d[i] == 0xFF and (d[i + 1] & 0xE6) == 0xE2):  # sincronia + Layer III
        i += 1
    if i >= limite:
        return None
    versao, ib, it, modo = (d[i + 1] >> 3) & 3, d[i + 2] >> 4, (d[i + 2] >> 2) & 3, d[i + 3] >> 6
    if versao == 1 or ib in (0, 15) or it == 3:
        return None
    taxa = _TAXAS_MP3[versao][it]
    amostras = 1152 if versao == 3 else 576
    lado = (32 if modo != 3 else 17) if versao == 3 else (17 if modo != 3 else 9)
    xing = d[i + 4 + lado:i + 4 + lado + 12]
    if xing[:4] in (b"Xing", b"Info") and struct.unpack(">I", xing[4:8])[0] & 1:
        return struct.unpack(">I", xing[8:12])[0] * amostras / taxa
    kbps = _BITRATES_MP3[1 if versao == 3 else 2][ib]
    return (len(d) - i) * 8 / (kbps * 1000) if kbps else None


def duracao_do_audio(dados: bytes) -> Optional[float]:
    """Segundos de áudio pelo cabeçalho (wav · m4a/mp4 · ogg/opus · webm · mp3), ou `None` quando ele não diz
    (o WebM do gravador do navegador costuma não trazer a duração — aí vale só o teto de 20 MB). Nunca lança."""
    for leitor in (_duracao_wav, _duracao_mp4, _duracao_ogg, _duracao_webm, _duracao_mp3):
        try:
            segundos = leitor(dados)
        except Exception:  # noqa: BLE001 — cabeçalho estranho: não se sabe a duração
            segundos = None
        if segundos is not None and segundos >= 0:
            return float(segundos)
    return None


# =====================================================================================================================
# planilha (xlsx · xls · csv) — o leitor TABULAR
# =====================================================================================================================
def _sem_acento(valor: Any) -> str:
    t = unicodedata.normalize("NFKD", str(valor or ""))
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


#: Coluna cujo NOME diz o que ela guarda → o marcador da 130-B na célula inteira (a máscara de texto não reconhece
#: uma data de nascimento solta numa célula, nem um CPF que o Excel guardou como número sem o zero da frente).
_COLUNAS_SENSIVEIS: Tuple[Tuple[re.Pattern, str], ...] = (
    (re.compile(r"\bcnpj\b"), "[CNPJ]"),
    (re.compile(r"\bcpf\b"), "[CPF]"),
    (re.compile(r"\brg\b|\bidentidade\b"), "[RG]"),
    (re.compile(r"\bnasc|\bnascimento\b|\bdt nasc|\bdata de nasc"), "[DATA]"),
    (re.compile(r"\btelefone\b|\bcelular\b|\bfone\b|\btel\b|\bwhatsapp\b"), "[TELEFONE]"),
    (re.compile(r"\be ?mail\b"), "[EMAIL]"),
    (re.compile(r"\bcep\b"), "[CEP]"),
    (re.compile(r"\bplaca\b"), "[PLACA]"),
    (re.compile(r"\bchassi\b"), "[CHASSI]"),
)
#: Colunas de identificador: nunca somadas (a soma de números de apólice não diz nada, e o mín/máx expõe um).
#: Pelo COMEÇO do nome: "Nº da apólice" e "Apólice" são identificadores; "Prêmio da apólice" é dinheiro.
_COLUNAS_DE_IDENTIFICADOR = re.compile(r"^(?:id|codigo|cod|ano|n|no|numero|num|matricula|apolice|proposta|endosso)\b")


def _marcador_da_coluna(cabecalho: str) -> Optional[str]:
    n = _sem_acento(cabecalho)
    for padrao, marcador in _COLUNAS_SENSIVEIS:
        if padrao.search(n):
            return marcador
    return None


def _formatar_numero(v: float) -> str:
    """Inteiro sem separador (o CPF que o Excel guardou como número cai na máscara de dígitos crus);
    decimal com vírgula e 2 casas, como o corretor lê."""
    if v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return _formatar_dinheiro(v)


def _formatar_dinheiro(v: float) -> str:
    """`1.234,56` — a coluna que tem centavos é somada e mostrada como o corretor lê dinheiro."""
    inteiro, _, frac = f"{abs(v):,.2f}".partition(".")
    return ("-" if v < 0 else "") + inteiro.replace(",", ".") + "," + frac


def _celula_em_texto(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "sim" if v else "não"
    if isinstance(v, (int, float)):
        return _formatar_numero(float(v))
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y") + ("" if v.time() == time(0, 0) else v.strftime(" %H:%M"))
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return " ".join(str(v).split())


_NUMERO_BR = re.compile(r"^-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?$")
_NUMERO_PONTO = re.compile(r"^-?\d+\.\d+$")


def _numero(v: Any) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "").replace(" ", "")
    if not s:
        return None
    try:
        if _NUMERO_BR.match(s):
            return float(s.replace(".", "").replace(",", "."))
        if _NUMERO_PONTO.match(s):
            return float(s)
    except ValueError:
        return None
    return None


def _celula_md(texto: str) -> str:
    """A célula como vai na tabela markdown: até 80 caracteres.

    🔴 Conserto da SPEC-132-A (red team B1): a célula longa é MASCARADA INTEIRA antes do corte. 📊 14/10: o corte vinha
    antes e um CPF que cruzava o caractere 77 saía com 9–10 dos 11 dígitos ("529.982.247-2...") — a regex do CPF e o
    Presidio não reconheciam mais o pedaço, e os 9 dígitos-base recalculam o CPF inteiro. Máscara primeiro, corte
    depois. ⛔ Máscara indisponível → nenhum pedaço da célula (fail-closed)."""
    t = texto.replace("|", "/").replace("\n", " ")
    if len(t) <= 80:
        return t
    mascarado = _mascarar(t)
    if not mascarado:
        return "[texto omitido]"
    t = mascarado.replace("|", "/").replace("\n", " ")
    return t if len(t) <= 80 else t[:77] + "..."


def _vazia(v: Any) -> bool:
    return v is None or v == "" or (isinstance(v, str) and not v.strip())


class _Orcamento:
    """O teto de CÉLULAS de um arquivo (todas as abas) e o PRAZO do turno (conferido a cada 💭 256 linhas)."""

    def __init__(self, celulas: int = CELULAS_POR_ARQUIVO, prazo: Optional[float] = None):
        self.restante = celulas
        self.prazo = prazo
        self._passos = 0

    @property
    def esgotado(self) -> bool:
        return self.restante <= 0

    def gastar(self, celulas: int) -> bool:
        """Desconta uma linha. `False` = o teto acabou (a linha NÃO é lida). Prazo vencido → levanta."""
        self._passos += 1
        if self.prazo is not None and self._passos % 256 == 0 and _relogio.monotonic() > self.prazo:
            raise PrazoDaLeituraEstourado()
        if self.restante <= 0:
            return False
        self.restante -= max(1, celulas)
        return True


class _Coluna:
    """A estatística de UMA coluna, em fluxo (sem guardar os valores)."""

    __slots__ = ("n", "soma", "minimo", "maximo", "decimal", "maior_abs", "invalida")

    def __init__(self):
        self.n, self.soma, self.minimo, self.maximo = 0, 0.0, None, None
        self.decimal, self.maior_abs, self.invalida = False, 0.0, False

    def somar(self, v: Any) -> None:
        if self.invalida:
            return
        x = _numero(v)
        if x is None:
            self.invalida = True  # um valor que não é número tira a coluna dos totais (a regra de antes)
            return
        self.n += 1
        self.soma += x
        self.minimo = x if self.minimo is None else min(self.minimo, x)
        self.maximo = x if self.maximo is None else max(self.maximo, x)
        self.decimal = self.decimal or x != int(x)
        self.maior_abs = max(self.maior_abs, abs(x))


def _uma_aba(nome_da_aba: str, linhas: Iterable[Any], orcamento: Any = None, *, colunas_alem: bool = False) -> str:
    """Uma aba em texto: cabeçalho, amostra de 💭 50 linhas × 30 colunas e os totais das colunas numéricas.

    🔴 Lê em FLUXO, linha a linha: só a amostra e uma estatística por coluna ficam em memória. Nenhuma linha é
    completada até a mais larga (era isso que fazia 14 KB de CSV pedirem 154 MB). Cada linha desconta do `orcamento`
    de células do arquivo; acabou → a contagem leva "+" e a leitura para.
    """
    if not isinstance(orcamento, _Orcamento):
        orcamento = _Orcamento()
    cab_bruto: Optional[List[Any]] = None
    largura_real = 0
    total = 0
    lidas = 0
    truncou = False
    mais_colunas = colunas_alem
    amostra: List[List[Any]] = []
    colunas: List[_Coluna] = []
    pular: List[bool] = []
    for ln in linhas:
        if lidas >= LINHAS_LIDAS_POR_ABA:
            truncou = True
            break
        lidas += 1
        n = len(ln)
        if not orcamento.gastar(n):
            truncou = True
            break
        fim = n
        while fim and _vazia(ln[fim - 1]):
            fim -= 1
        if not fim:
            continue
        if fim > COLUNAS_LIDAS:
            mais_colunas = True
        celulas = list(ln[:min(fim, COLUNAS_LIDAS)])
        largura_real = max(largura_real, fim)
        if cab_bruto is None:
            cab_bruto = celulas
            nomes = [_sem_acento(_celula_em_texto(c)) for c in cab_bruto]
            pular = [bool(_marcador_da_coluna(_celula_em_texto(c)) or _COLUNAS_DE_IDENTIFICADOR.search(nm))
                     for c, nm in zip(cab_bruto, nomes)]
            continue
        total += 1
        if len(amostra) < LINHAS_DA_AMOSTRA:
            amostra.append(celulas[:COLUNAS_DA_AMOSTRA])
        for i, v in enumerate(celulas):
            if _vazia(v) or (i < len(pular) and pular[i]):
                continue
            while len(colunas) <= i:
                colunas.append(_Coluna())
            colunas[i].somar(v)

    if cab_bruto is None:
        if truncou:
            return f'## Aba "{nome_da_aba}" — não lida (a planilha é grande demais para ler inteira)'
        return f'## Aba "{nome_da_aba}" — vazia'
    largura = min(largura_real, COLUNAS_LIDAS)
    cab_bruto = cab_bruto + [None] * (largura - len(cab_bruto))
    cabecalho = [(_celula_em_texto(c) or f"coluna {i + 1}") for i, c in enumerate(cab_bruto)]
    marcadores = [_marcador_da_coluna(c) for c in cabecalho]

    def _celula(i: int, ln: List[Any]) -> str:
        v = ln[i] if i < len(ln) else None
        if marcadores[i] and not _vazia(v):
            return marcadores[i]
        return _celula_em_texto(v)

    out = [f'## Aba "{nome_da_aba}" — {total}{"+" if truncou else ""} linhas de dados × {largura_real} colunas',
           "Colunas: " + ", ".join(cabecalho)
           + (f" (e mais {largura_real - largura})" if largura_real > largura else "")]
    if truncou:
        out.append("(a planilha é grande: li só o começo dela)")
    if mais_colunas:
        out.append(f"(li só as primeiras {COLUNAS_LIDAS} colunas)")
    mostradas = list(range(min(largura, COLUNAS_DA_AMOSTRA)))
    if total:
        out.append("| " + " | ".join(_celula_md(cabecalho[i]) for i in mostradas) + " |")
        out.append("|" + "---|" * len(mostradas))
        for ln in amostra:
            out.append("| " + " | ".join(_celula_md(_celula(i, ln)) for i in mostradas) + " |")
        if total > LINHAS_DA_AMOSTRA:
            out.append(f"(mostrando {LINHAS_DA_AMOSTRA} de {total} linhas)")
        if largura_real > COLUNAS_DA_AMOSTRA:
            out.append(f"(mostrando {COLUNAS_DA_AMOSTRA} de {largura_real} colunas)")
    totais = []
    for i, c in enumerate(colunas[:largura]):
        if c.invalida or not c.n or marcadores[i] or _COLUNAS_DE_IDENTIFICADOR.search(_sem_acento(cabecalho[i])):
            continue
        # 8+ dígitos inteiros parece identificador (CPF sem pontuação, telefone, número de documento): não se soma
        if not c.decimal and c.maior_abs >= 1e7:
            continue
        fmt = _formatar_dinheiro if c.decimal else _formatar_numero
        totais.append(f"- {cabecalho[i]}: soma {fmt(c.soma)} · mín {fmt(c.minimo)} · "
                      f"máx {fmt(c.maximo)} ({c.n} valores)")
    if totais:
        out.append("Totais das colunas numéricas:")
        out.extend(totais)
    return "\n".join(out)


# ---------------------------------------------------------------------------------------------------------------------
# o pacote Office (zip): descomprimido com teto, e a linha forjada recusada ANTES do openpyxl
# ---------------------------------------------------------------------------------------------------------------------
_MARCA_DO_XML = re.compile(rb"<(?:[A-Za-z_][\w.-]*:)?(row|c)[\s>/]")
_ATRIBUTO_R = re.compile(rb'\sr="(\d+)"')


def _conferir_o_pacote(dados: bytes, *, por_parte: int) -> None:
    """Levanta `ArquivoGrandeDemais` se o zip descomprime além dos tetos (a medida DECLARADA é a que o `zipfile` lê:
    ele nunca entrega mais bytes que o `file_size`). Zip ilegível passa — o leitor diz que está corrompido."""
    try:
        pacote = zipfile.ZipFile(io.BytesIO(dados))
    except Exception:  # noqa: BLE001
        return
    with pacote:
        total = 0
        for info in pacote.infolist():
            total += info.file_size
            teto = DESCOMPRIMIDO_DAS_STRINGS if info.filename.endswith("sharedStrings.xml") else por_parte
            if info.file_size > teto or total > DESCOMPRIMIDO_TOTAL:
                raise ArquivoGrandeDemais(info.filename)
            if info.file_size > 10 * 1024 * 1024 and info.file_size > RAZAO_SUSPEITA * max(1, info.compress_size):
                raise ArquivoGrandeDemais(info.filename)


def _ultima_linha_no_orcamento(pacote: zipfile.ZipFile, parte: str, celulas: int) -> Optional[int]:
    """Varre o XML da aba em pedaços (memória constante) contando `<c>` por `<row>`.

    Levanta `ArquivoGrandeDemais` se uma linha passa de 💭 20.000 células (o openpyxl monta a linha INTEIRA na memória
    antes de devolvê-la — 10 milhões de `<c/>` cabem em 40 KB de zip). Devolve o número da última linha que cabe no
    orçamento de células (o `max_row` do openpyxl), ou `None` se a aba inteira cabe."""
    vistas, na_linha, numero = 0, 0, 0
    resto = b""
    with pacote.open(parte) as fonte:
        while True:
            pedaco = fonte.read(1 << 20)
            buf = resto + pedaco
            corte = len(buf) if not pedaco else max(0, len(buf) - 64)
            for m in _MARCA_DO_XML.finditer(buf):
                if m.start() >= corte:
                    break
                if m.group(1) == b"row":
                    fecha = buf.find(b">", m.start())
                    r = _ATRIBUTO_R.search(buf, m.start(), fecha) if fecha != -1 else None
                    numero = int(r.group(1)) if r else numero + 1
                    na_linha = 0
                else:
                    na_linha += 1
                    vistas += 1
                    if na_linha > CELULAS_POR_LINHA_DO_XML:
                        raise ArquivoGrandeDemais(parte)
                    if vistas > celulas:
                        return max(1, numero)
            if not pedaco:
                return None
            resto = buf[corte:]


def _abas_do_xlsx(dados: bytes, orcamento: _Orcamento) -> List[str]:
    _conferir_o_pacote(dados, por_parte=DESCOMPRIMIDO_POR_PARTE_DA_PLANILHA)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    try:
        pacote = zipfile.ZipFile(io.BytesIO(dados))
        saida = []
        for ws in wb.worksheets[:ABAS_LIDAS]:
            parte = str(getattr(ws, "_worksheet_path", "") or "").lstrip("/")
            max_row = None
            if parte and parte in pacote.namelist():
                max_row = _ultima_linha_no_orcamento(pacote, parte, max(1, orcamento.restante))
            try:
                maxc = int(ws.max_column) if ws.max_column else None
            except Exception:  # noqa: BLE001 — aba sem dimensão: a linha vem só até a última célula
                maxc = None
            limite = min(maxc, COLUNAS_LIDAS) if maxc else None
            linhas = ws.iter_rows(values_only=True, max_col=limite, max_row=max_row)
            saida.append(_uma_aba(str(ws.title), linhas, orcamento, colunas_alem=bool(maxc and maxc > COLUNAS_LIDAS)))
        return saida
    finally:
        wb.close()


def _abas_do_xls(dados: bytes, orcamento: _Orcamento) -> List[str]:
    import xlrd

    wb = xlrd.open_workbook(file_contents=dados, on_demand=True)
    try:
        saida = []
        for idx in range(min(wb.nsheets, ABAS_LIDAS)):
            sh = wb.sheet_by_index(idx)

            def _linhas(sh=sh):
                for r in range(sh.nrows):
                    ln = []
                    for c in range(min(sh.row_len(r), COLUNAS_LIDAS)):
                        cel = sh.cell(r, c)
                        if cel.ctype == xlrd.XL_CELL_DATE:
                            try:
                                ln.append(xlrd.xldate_as_datetime(cel.value, wb.datemode))
                            except Exception:  # noqa: BLE001
                                ln.append(cel.value)
                        elif cel.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            ln.append(None)
                        elif cel.ctype == xlrd.XL_CELL_BOOLEAN:
                            ln.append(bool(cel.value))
                        else:
                            ln.append(cel.value)
                    yield ln

            saida.append(_uma_aba(str(sh.name), _linhas(), orcamento, colunas_alem=sh.ncols > COLUNAS_LIDAS))
            wb.unload_sheet(idx)
        return saida
    finally:
        wb.release_resources()


def _decodificar_texto(dados: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return dados.decode("utf-8", errors="replace")


def _codificacao_do_texto(dados: bytes) -> str:
    """A mesma escolha de `_decodificar_texto` (utf-8 com BOM → cp1252 → utf-8 com troca), sem montar o texto inteiro."""
    for codificacao in ("utf-8-sig", "cp1252"):
        decodificador = codecs.getincrementaldecoder(codificacao)()
        try:
            for i in range(0, len(dados), 1 << 20):
                decodificador.decode(dados[i:i + (1 << 20)], final=False)
            decodificador.decode(b"", final=True)
            return codificacao
        except UnicodeDecodeError:
            continue
    return "utf-8"


def _abas_do_csv(dados: bytes, orcamento: _Orcamento) -> List[str]:
    codificacao = _codificacao_do_texto(dados)
    amostra = dados[:8192].decode(codificacao, errors="replace")
    try:
        separador = csv.Sniffer().sniff(amostra, delimiters=";,\t|").delimiter
    except csv.Error:
        primeira = amostra.splitlines()[0] if amostra.splitlines() else ""
        separador = max(";,\t|", key=primeira.count)
    texto = io.TextIOWrapper(io.BytesIO(dados), encoding=codificacao, errors="replace", newline="")

    def _fisicas():
        for linha in texto:
            yield linha if len(linha) <= CARACTERES_POR_LINHA_DO_CSV else linha[:CARACTERES_POR_LINHA_DO_CSV] + "\n"

    return [_uma_aba("CSV", csv.reader(_fisicas(), delimiter=separador), orcamento)]


def tabela_da_planilha(dados: bytes, tipo: str, *, prazo: Optional[float] = None) -> str:
    """O texto da planilha (cabeçalho, amostra, totais por aba). ⚠️ Só sai deste módulo pela máscara.

    `prazo` (relógio monotônico): passado ele, a leitura PARA e levanta `PrazoDaLeituraEstourado` — a thread volta ao
    poço em vez de seguir lendo o que ninguém vai mais usar."""
    leitor = {"xlsx": _abas_do_xlsx, "xls": _abas_do_xls, "csv": _abas_do_csv}[tipo]
    abas = leitor(dados, _Orcamento(prazo=prazo))
    cabeca = f"Planilha com {len(abas)} aba(s)." if tipo != "csv" else "Arquivo CSV."
    return "\n\n".join([cabeca] + abas)


# =====================================================================================================================
# Word e PowerPoint
# =====================================================================================================================
def _xml_em_texto(xml: str, marca_de_texto: str, marca_de_paragrafo: str) -> str:
    import html

    paragrafos = []
    for p in re.split(rf"</{marca_de_paragrafo}>", xml):
        pedacos = re.findall(rf"<{marca_de_texto}(?:\s[^>]*)?>([^<]*)</{marca_de_texto}>", p)
        linha = html.unescape("".join(pedacos)).strip()
        if linha:
            paragrafos.append(linha)
    return "\n".join(paragrafos)


def _texto_do_docx(dados: bytes) -> str:
    try:
        import docx  # python-docx

        doc = docx.Document(io.BytesIO(dados))
        partes = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for tabela in doc.tables:
            for linha in tabela.rows:
                celulas = [" ".join(c.text.split()) for c in linha.cells]
                if any(celulas):
                    partes.append(" | ".join(celulas))
        return "\n".join(partes)
    except ImportError:
        pass
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        return _xml_em_texto(z.read("word/document.xml").decode("utf-8", errors="replace"), "w:t", "w:p")


def _texto_do_pptx(dados: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[-1]))
        partes = []
        for i, nome in enumerate(slides, 1):
            texto = _xml_em_texto(z.read(nome).decode("utf-8", errors="replace"), "a:t", "a:p")
            if texto:
                partes.append(f"[slide {i}]\n{texto}")
        return "\n\n".join(partes)


def _docling_instalado() -> bool:
    import importlib.util

    try:
        return importlib.util.find_spec("docling") is not None
    except Exception:  # noqa: BLE001
        return False


# =====================================================================================================================
# um anexo
# =====================================================================================================================
@dataclass
class _Leitura:
    tipo: str = ""
    bloco: str = ""
    imagem: str = ""
    descricao: str = ""
    transcricao: str = ""
    aviso: str = ""
    lido: bool = False


def _bloco(nome: str, corpo: str) -> str:
    return f"[DOCUMENTO ANEXADO: {nome}]\n{corpo}\n[FIM DO DOCUMENTO]"


async def _texto_mascarado(bruto: str) -> Optional[str]:
    if not str(bruto or "").strip():
        return None
    mascarado = _mascarar(bruto)
    if not mascarado or not mascarado.strip():
        return None
    return _truncar(mascarado)


async def _ler_um(company_id: str, anexo: Dict[str, Any], nome: str, tipo: str,
                  ao_progredir: Optional[Callable[[str], Awaitable[None]]], *,
                  prazo: Optional[float] = None) -> _Leitura:
    r = _Leitura(tipo=tipo)
    bucket = str(anexo.get("bucket") or "")
    path = anexo.get("path")
    mime = str(anexo.get("mime") or "").split(";")[0].strip().lower()
    nome_cru = str(anexo.get("nome") or "")

    if bucket not in BUCKETS_DO_CHAT or not path_da_corretora(path, company_id):
        r.aviso = f"O anexo «{nome}» não pertence a esta corretora e não foi lido."
        return r
    declarado = anexo.get("bytes")
    if isinstance(declarado, (int, float)) and declarado > LIMITE_DE_BYTES:
        r.aviso = f"O arquivo «{nome}» passa de 20 MB e não foi lido."
        return r
    if not tipo:
        r.aviso = f"Ainda não sei ler o formato de «{nome}»."
        return r

    if ao_progredir is not None:
        frase = {"imagem": f"Olhando a foto {nome}…", "audio": f"Ouvindo o áudio {nome}…"}.get(tipo, f"Lendo {nome}…")
        try:
            await ao_progredir(frase)
        except Exception:  # noqa: BLE001 — a linha de atividade nunca derruba a leitura
            pass

    try:
        dados = await _baixar_do_storage(bucket, str(path))
    except Exception as e:  # noqa: BLE001
        logger.warning("[ANEXOS] download falhou (%s)", type(e).__name__)
        dados = None
    if not dados:
        r.aviso = f"Não encontrei o arquivo «{nome}». Pode mandar de novo?"
        return r
    if len(dados) > LIMITE_DE_BYTES:
        r.aviso = f"O arquivo «{nome}» passa de 20 MB e não foi lido."
        return r

    if tipo == "imagem":
        jpeg, tipo_img, aviso = await _no_poco(_imagem_para_jpeg, dados, mime, nome_cru)
        if not jpeg:
            r.aviso = (aviso or "Não consegui abrir a foto «{nome}».").replace("{nome}", nome)
            return r
        from app.services.vision_service import data_url_da_imagem

        r.imagem = data_url_da_imagem(jpeg, tipo_img)
        r.lido = True
        try:
            desc = await _descrever_imagem(r.imagem, company_id)
        except Exception as e:  # noqa: BLE001
            logger.warning("[ANEXOS] descrição da foto falhou (%s)", type(e).__name__)
            desc = None
        if desc:
            mascarada = _mascarar(desc)
            if mascarada and mascarada.strip():
                r.descricao = mascarada.strip()
        return r

    if tipo == "pdf":
        from app.services import vision_service

        corpo = await vision_service.conteudo_do_anexo(
            dados, nome_cru or "documento.pdf", company_id=company_id, audiencia="corretor",
            orcamento_s=max(5.0, ORCAMENTO_TOTAL_S - 10.0))
        if not corpo:
            r.aviso = (f"Não consegui ler o PDF «{nome}». Se ele for uma foto digitalizada, mande as páginas como "
                       "imagem — ou cole aqui o trecho que importa.")
            return r
        r.bloco, r.lido = _bloco(nome, corpo), True
        return r

    if tipo in ("xlsx", "xls", "csv"):
        try:
            bruto = await _no_poco(tabela_da_planilha, dados, tipo, prazo=prazo)
        except ImportError:
            r.aviso = f"Ainda não consigo abrir planilhas .{tipo} aqui. Pode salvar «{nome}» como .xlsx ou .csv?"
            return r
        except ArquivoGrandeDemais:
            r.aviso = (f"A planilha «{nome}» é grande demais para eu ler aqui. Pode mandar só a aba que importa — ou "
                       "salvar como .csv?")
            return r
        except PrazoDaLeituraEstourado:
            r.aviso = f"Demorei demais para ler «{nome}». Pode mandar de novo, ou só o trecho que importa?"
            return r
        except Exception as e:  # noqa: BLE001
            logger.warning("[ANEXOS] planilha não abriu (%s)", type(e).__name__)
            r.aviso = f"Não consegui abrir a planilha «{nome}». Ela pode estar protegida por senha ou corrompida."
            return r
        corpo = await _texto_mascarado(bruto)
        if not corpo:
            r.aviso = f"Não consegui ler a planilha «{nome}» com segurança agora."
            return r
        r.bloco, r.lido = _bloco(nome, corpo), True
        return r

    if tipo in ("docx", "pptx"):
        corpo = None
        try:
            # o pacote descomprimido com teto ANTES de qualquer leitor (python-docx monta o XML inteiro na memória)
            _conferir_o_pacote(dados, por_parte=DESCOMPRIMIDO_POR_PARTE_DO_DOCUMENTO)
        except ArquivoGrandeDemais:
            r.aviso = f"O arquivo «{nome}» é grande demais para eu ler aqui. Pode colar aqui o trecho que importa?"
            return r
        if _docling_instalado():
            from app.services import vision_service

            corpo = await vision_service.texto_mascarado_do_documento(dados, nome_cru or f"documento.{tipo}")
        if not corpo:
            try:
                bruto = await _no_poco(_texto_do_docx if tipo == "docx" else _texto_do_pptx, dados)
            except Exception as e:  # noqa: BLE001
                logger.warning("[ANEXOS] %s não abriu (%s)", tipo, type(e).__name__)
                bruto = ""
            corpo = await _texto_mascarado(bruto)
        if not corpo:
            r.aviso = f"Não consegui ler o arquivo «{nome}». Pode colar aqui o trecho que importa?"
            return r
        r.bloco, r.lido = _bloco(nome, corpo), True
        return r

    if tipo == "texto":
        corpo = await _texto_mascarado(_decodificar_texto(dados))
        if not corpo:
            r.aviso = f"O arquivo «{nome}» está vazio ou não consegui lê-lo."
            return r
        r.bloco, r.lido = _bloco(nome, corpo), True
        return r

    if tipo == "audio":
        # 🔴 red team P6: o arquivo de 20 MB (~1 h em opus) ia inteiro à transcrição, cobrado pelo ledger.
        duracao = duracao_do_audio(dados)
        if duracao is not None and duracao > SEGUNDOS_MAXIMOS_DO_AUDIO:
            r.aviso = AVISO_AUDIO_LONGO.replace("{nome}", nome)
            return r
        try:
            texto = await _transcrever_audio(dados, mime, nome_cru, company_id)
        except Exception as e:  # noqa: BLE001
            logger.warning("[ANEXOS] transcrição falhou (%s)", type(e).__name__)
            texto = ""
        if not str(texto or "").strip():
            r.aviso = f"Não consegui ouvir o áudio «{nome}». Pode escrever em poucas palavras o que ele diz?"
            return r
        mascarado = _mascarar(texto)
        if not mascarado or not mascarado.strip():
            r.aviso = f"Não consegui ouvir o áudio «{nome}» com segurança agora. Pode escrever o que ele diz?"
            return r
        r.transcricao, r.lido = mascarado.strip(), True
        return r

    r.aviso = f"Ainda não sei ler o formato de «{nome}»."
    return r


# =====================================================================================================================
# o contrato (A2 → A3)
# =====================================================================================================================
async def preparar(
    company_id: str,
    anexos: List[Dict[str, Any]],
    *,
    ao_progredir: Optional[Callable[[str], Awaitable[None]]] = None,
) -> ResultadoDosAnexos:
    """Lê os anexos de UM turno do chat, em paralelo, com orçamento total de 💭 45 s. Nunca lança.

    `anexos`: a lista que o BFF repassou — `[{bucket, path, nome, mime, bytes}]`, cada path já conferido contra a
    corretora da sessão (e conferido de novo aqui). `ao_progredir("Lendo apolice.pdf…")` é chamado por arquivo,
    antes do download — é a linha de atividade da tela (evento `activity`).
    """
    resultado = ResultadoDosAnexos()
    company_id = str(company_id or "").strip()
    if company_id.lower() == "none":
        company_id = ""
    lista = [a for a in (anexos or []) if isinstance(a, dict)]
    if not lista:
        return resultado
    if len(lista) > MAX_ANEXOS:
        resultado.avisos.append(f"Li só os {MAX_ANEXOS} primeiros arquivos desta mensagem.")
        lista = lista[:MAX_ANEXOS]

    preparados = []
    for a in lista:
        nome = _nome_para_o_modelo(str(a.get("nome") or "") or str(a.get("path") or "").rsplit("/", 1)[-1])
        tipo = tipo_do_anexo(str(a.get("mime") or ""), str(a.get("nome") or ""), str(a.get("path") or ""))
        preparados.append((a, nome, tipo))

    prazo = _relogio.monotonic() + ORCAMENTO_TOTAL_S
    tarefas = [asyncio.ensure_future(_ler_um(company_id, a, nome, tipo, ao_progredir, prazo=prazo))
               for a, nome, tipo in preparados]
    try:
        _feitas, pendentes = await asyncio.wait(tarefas, timeout=ORCAMENTO_TOTAL_S)
    except Exception as e:  # noqa: BLE001
        logger.error("[ANEXOS] espera falhou (%s)", type(e).__name__)
        pendentes = {t for t in tarefas if not t.done()}
    for t in pendentes:
        t.cancel()

    blocos: List[str] = []
    for tarefa, (a, nome, tipo) in zip(tarefas, preparados):
        if tarefa in pendentes:
            leitura = _Leitura(tipo=tipo, aviso=f"Demorei demais para ler «{nome}». Pode mandar de novo, ou só o "
                                                 "trecho que importa?")
        else:
            try:
                leitura = tarefa.result()
            except Exception as e:  # noqa: BLE001 — um arquivo que estoura não derruba os outros
                logger.error("[ANEXOS] leitura estourou (%s)", type(e).__name__)
                leitura = _Leitura(tipo=tipo, aviso=f"Não consegui ler «{nome}». Pode mandar de novo?")
        if leitura.bloco:
            blocos.append(leitura.bloco)
        if leitura.imagem:
            resultado.imagens.append(leitura.imagem)
        if leitura.descricao:
            resultado.descricoes.append(leitura.descricao)
            resultado._nomes_das_descricoes.append(nome)
        if leitura.transcricao:
            resultado.transcricoes.append(leitura.transcricao)
            resultado._nomes_das_transcricoes.append(nome)
        if leitura.aviso:
            resultado.avisos.append(leitura.aviso)
        resultado.lidos.append({
            "nome": nome,
            "mime": str(a.get("mime") or "")[:100],
            "bytes": a.get("bytes") if isinstance(a.get("bytes"), (int, float)) else None,
            "path": a.get("path") if isinstance(a.get("path"), str) else None,
            "bucket": a.get("bucket") if isinstance(a.get("bucket"), str) else None,
            "tipo": tipo or "desconhecido",
            "lido": bool(leitura.lido),
        })

    texto = "\n\n".join(blocos)
    if len(texto) > TETO_DO_TEXTO:
        texto = texto[:TETO_DO_TEXTO] + "\n\n[ANEXOS TRUNCADOS — o conteúdo acima cobre o início dos arquivos]"
        resultado.avisos.append("Os arquivos eram longos demais para ler inteiros; li o começo de cada um.")
    resultado.texto = texto

    contagem: Dict[str, int] = {}
    for item in resultado.lidos:
        chave = f"{item['tipo']}:{'ok' if item['lido'] else 'aviso'}"
        contagem[chave] = contagem.get(chave, 0) + 1
    logger.info("[ANEXOS] %d anexo(s) · %s", len(resultado.lidos), contagem)
    return resultado
